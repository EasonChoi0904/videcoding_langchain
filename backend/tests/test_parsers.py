"""docx / html 解析器测试(纯本地文件,无外网依赖)。"""
from pathlib import Path

import docx

from app.rag.splitter import parse_document, split_pieces


class TestDocxParse:
    def test_paragraphs_and_table_rows(self, tmp_path: Path):
        f = tmp_path / "说明.docx"
        d = docx.Document()
        d.add_heading("商品说明书", level=1)
        d.add_paragraph("支持 90W 快充,电池 5400mAh。")
        tbl = d.add_table(rows=2, cols=2)
        tbl.rows[0].cells[0].text = "屏幕"
        tbl.rows[0].cells[1].text = "6.7 英寸"
        tbl.rows[1].cells[0].text = "价格"
        tbl.rows[1].cells[1].text = "4499 元"
        d.save(str(f))

        pieces = parse_document("docx", f)
        texts = [t for t, _ in pieces]
        joined = "".join(texts)
        # 段落文本进入知识库
        assert "5400mAh" in joined
        # 表格行被拼成"单元行"保留(单元格内容出现)
        assert "6.7 英寸" in joined
        assert "4499 元" in joined
        # 表格行带 type 标记
        assert any(meta.get("type") == "table-row" for _, meta in pieces)

        splits = split_pieces(pieces)
        assert len(splits) >= 1
        # 分块后无空块
        assert all(t.strip() for t, _, _ in splits)


class TestHtmlParse:
    HTML = """<!DOCTYPE html>
<html><head><title>商品页 - 冰箱501L</title></head>
<body>
<nav>首页 | 关于 | 联系</nav>
<div class="product">
  <h1>海尔零嵌冰箱 501L</h1>
  <p>总容积 501 升,一级能效,风冷无霜。</p>
  <p>支持零嵌安装,两侧各留 2cm 即可散热。</p>
</div>
<script>alert("噪声脚本")</script>
<style>.noise{color:red}</style>
<footer>版权声明 2026</footer>
</body></html>
"""

    def test_noise_removed_and_body_kept(self, tmp_path: Path):
        f = tmp_path / "page.html"
        f.write_text(self.HTML, encoding="utf-8")
        pieces = parse_document("html", f)
        texts = [t for t, _ in pieces]
        joined = " ".join(texts)

        # 标题进知识库
        assert "商品页" in joined
        # 正文保留
        assert "501" in joined and "风冷无霜" in joined
        # 噪声剔除:脚本内容 / 样式 / 版权页脚不应进入
        assert "alert" not in joined
        assert ".noise" not in joined
        assert "版权声明" not in joined

    def test_cjk_gbk_file_safe(self, tmp_path: Path):
        """用 GBK 编码的中文 HTML 也应能解析(编码容错)。"""
        f = tmp_path / "gbk.html"
        f.write_bytes("<html><title>冰箱说明</title><body><p>容量 501 升</p></body></html>".encode("gbk"))
        pieces = parse_document("html", f)
        joined = " ".join(t for t, _ in pieces)
        assert "容量 501 升" in joined
